import os
import csv
import io
from typing import List, Dict, Any
from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl

class ParsedSection:
    def __init__(self, content: str, page_number: int = 1, section_title: str = "Main"):
        self.content = content
        self.page_number = page_number
        self.section_title = section_title

class DocumentParser:
    @staticmethod
    def parse_file(file_path: str, file_type: str) -> List[ParsedSection]:
        ext = file_type.lower().lstrip(".")
        if ext == "pdf":
            return DocumentParser._parse_pdf(file_path)
        elif ext in ["docx", "doc"]:
            return DocumentParser._parse_docx(file_path)
        elif ext in ["xlsx", "xls"]:
            return DocumentParser._parse_xlsx(file_path)
        elif ext == "csv":
            return DocumentParser._parse_csv(file_path)
        elif ext in ["txt", "md", "markdown", "json"]:
            return DocumentParser._parse_text(file_path)
        else:
            # Fallback text
            return DocumentParser._parse_text(file_path)

    @staticmethod
    def _parse_pdf(file_path: str) -> List[ParsedSection]:
        sections: List[ParsedSection] = []
        reader = PdfReader(file_path)
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                sections.append(
                    ParsedSection(
                        content=text.strip(),
                        page_number=idx + 1,
                        section_title=f"Page {idx + 1}"
                    )
                )
        return sections

    @staticmethod
    def _parse_docx(file_path: str) -> List[ParsedSection]:
        sections: List[ParsedSection] = []
        doc = DocxDocument(file_path)
        current_title = "Document Body"
        buffer: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style.name.startswith("Heading"):
                if buffer:
                    sections.append(ParsedSection(content="\n".join(buffer), section_title=current_title))
                    buffer = []
                current_title = text
            buffer.append(text)

        if buffer:
            sections.append(ParsedSection(content="\n".join(buffer), section_title=current_title))
        return sections

    @staticmethod
    def _parse_xlsx(file_path: str) -> List[ParsedSection]:
        sections: List[ParsedSection] = []
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            lines: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                row_str = " | ".join(str(cell) for cell in row if cell is not None)
                if row_str.strip():
                    lines.append(row_str)
            if lines:
                sections.append(
                    ParsedSection(
                        content="\n".join(lines[:1000]),  # Limit sample size per sheet for RAG
                        section_title=f"Sheet: {sheet_name}"
                    )
                )
        return sections

    @staticmethod
    def _parse_csv(file_path: str) -> List[ParsedSection]:
        sections: List[ParsedSection] = []
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            lines = [" | ".join(row) for row in reader if row]
            if lines:
                sections.append(
                    ParsedSection(
                        content="\n".join(lines[:1000]),
                        section_title="CSV Data"
                    )
                )
        return sections

    @staticmethod
    def _parse_text(file_path: str) -> List[ParsedSection]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return [ParsedSection(content=text, section_title="Content")] if text.strip() else []
