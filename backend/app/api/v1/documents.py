"""Document ingestion and management routes."""

import csv
import io
import os
import uuid

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user
from app.core.exceptions import NotFoundException
from app.db.session import get_db, AsyncSessionLocal
from app.models.user import User
from app.repositories.ai_settings_repository import AISettingsRepository
from app.repositories.document_repository import DocumentRepository
from app.schemas.document import DocumentCreate, DocumentResponse, DocumentUpdate
from app.services.rag_service import RAGService

router = APIRouter(prefix="/documents", tags=["Documents"])


def _extract_text(filename: str, content: bytes) -> str:
    """Extract plain text from uploaded file bytes based on extension."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        import pypdf  # noqa: PLC0415

        reader = pypdf.PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext in (".xlsx", ".xls"):
        import openpyxl  # noqa: PLC0415

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append("\t".join("" if v is None else str(v) for v in row))
        return "\n".join(rows)

    if ext == ".docx":
        import docx  # noqa: PLC0415

        doc = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == ".csv":
        text = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        return "\n".join("\t".join(row) for row in reader)

    # Plain text / fallback
    return content.decode("utf-8", errors="replace")


async def _process_document_background(
    document_id: str,
    content: str,
    title: str,
    user_id: uuid.UUID,
) -> None:
    """Background task: chunk, embed and store a document after fast upload."""
    import logging  # noqa: PLC0415

    logger = logging.getLogger(__name__)
    async with AsyncSessionLocal() as db:
        try:
            ai = await AISettingsRepository(db).get_by_user(user_id)
            svc = RAGService(db, ai=ai)
            payload = DocumentCreate(title=title, content=content)
            await svc.ingest_document(payload, document_id=document_id)
            await db.commit()
        except Exception as exc:
            logger.exception("Background embedding failed for document %s: %s", document_id, exc)
            await db.rollback()
            # Write the failed status in a fresh transaction
            try:
                doc_repo = DocumentRepository(db)
                await doc_repo.set_embed_status(
                    uuid.UUID(document_id), "failed", str(exc)
                )
                await db.commit()
            except Exception:
                pass


@router.post("/upload", response_model=dict, status_code=status.HTTP_201_CREATED)
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """Upload a file immediately, then chunk+embed in the background."""
    allowed = {".pdf", ".docx", ".xlsx", ".xls", ".csv", ".txt"}
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in allowed:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(allowed))}",
        )
    raw = await file.read()
    max_bytes = 10 * 1024 * 1024  # 10 MB
    if len(raw) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File too large ({len(raw) // 1024 // 1024} MB). Maximum allowed size is 10 MB.",
        )
    try:
        content = _extract_text(file.filename or "file", raw)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Failed to parse file: {exc}",
        ) from exc

    title = os.path.splitext(file.filename or "Untitled")[0]

    # Save document record immediately with status=pending
    doc_repo = DocumentRepository(db)
    payload = DocumentCreate(title=title, content=content)
    document = await doc_repo.create(payload)
    await db.commit()  # Commit NOW — background task runs before get_db auto-commit
    document_id = str(document.id)

    # Schedule background processing
    background_tasks.add_task(
        _process_document_background,
        document_id=document_id,
        content=content,
        title=title,
        user_id=current_user.id,
    )

    return {"document_id": document_id, "chunks_created": 0, "embed_status": "pending"}

@router.post("/", response_model=dict, status_code=status.HTTP_201_CREATED)
async def ingest_document(
    payload: DocumentCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """Ingest a document into the RAG pipeline (chunk + embed + store)."""
    ai = await AISettingsRepository(db).get_by_user(current_user.id)
    svc = RAGService(db, ai=ai)
    return await svc.ingest_document(payload)


@router.get("/", response_model=list[DocumentResponse])
async def list_documents(
    limit: int = 50,
    offset: int = 0,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
) -> list[DocumentResponse]:
    """Return a paginated list of all ingested documents."""
    repo = DocumentRepository(db)
    documents = await repo.list_all(limit=limit, offset=offset)
    return [DocumentResponse.model_validate(doc) for doc in documents]


@router.get("/{document_id}", response_model=DocumentResponse)
async def get_document(
    document_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
) -> DocumentResponse:
    """Fetch a single document by UUID."""
    repo = DocumentRepository(db)
    document = await repo.get_by_id(document_id)
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return DocumentResponse.model_validate(document)


@router.patch("/{document_id}", response_model=DocumentResponse)
async def update_document(
    document_id: uuid.UUID,
    payload: DocumentUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
) -> DocumentResponse:
    """Update document metadata (title, source)."""
    repo = DocumentRepository(db)
    document = await repo.get_by_id(document_id)
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    updated = await repo.update(document, payload)
    return DocumentResponse.model_validate(updated)


@router.delete("/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    document_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
) -> None:
    """Delete a document and all its associated chunks."""
    repo = DocumentRepository(db)
    document = await repo.get_by_id(document_id)
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    await repo.delete(document)
